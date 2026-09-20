"""Tests for src/export.py -- flattening + PSV/XLSX rendering.

These use synthetic fixture data (control-flow / file-format testing only,
per project convention); they prove nothing about real Booking.com
retrieval and don't claim to.
"""

import io

from openpyxl import load_workbook

from src.export import EXPORT_COLUMNS, flatten_results, rows_to_psv, rows_to_xlsx_bytes

SAMPLE_RESULTS = [
    {
        "hotel": {
            "hotel_id": "hotel-a",
            "hotel_name": "Hotel A | Test",
            "booking_url": "https://www.booking.com/hotel/xx/hotel-a.html",
            "address": "1 Test Street",
            "city": "Testville",
            "country": "TC",
            "property_type": "Hotel",
            "amenities": ["Free Wifi", "Pool"],
            "rating_score": 8.7,
            "rating_label": "Wonderful",
            "review_count": 1234,
            "availability": "Only 1 left",
        },
        "units": [
            {
                "unit_name": "Double Room",
                "beds_text": "1 double bed",
                "size_m2": 20.0,
                "plan_name": "Breakfast included",
                "cancellation": "Free cancellation",
                "price_amount": 100.0,
                "price_currency": "USD",
                "amenities": ["TV", "Air conditioning"],
                "images": ["https://example.com/room.jpg"],
            }
        ],
    },
    {
        "hotel": {
            "hotel_id": "hotel-b",
            "hotel_name": "Hotel B",
            "booking_url": "https://www.booking.com/hotel/xx/hotel-b.html",
            "city": "Otherville",
            "country": "TC",
            "amenities": [],
        },
        "units": [],
    },
]


def test_flatten_basic_row_count():
    rows = flatten_results(SAMPLE_RESULTS)
    # 1 unit for hotel A, 1 blank-room row for hotel B (no units)
    assert len(rows) == 2


def test_flatten_missing_fields_are_blank_not_fabricated():
    rows = flatten_results(SAMPLE_RESULTS)
    row_b = next(r for r in rows if r["Hotel Name"] == "Hotel B")
    assert row_b["Rating"] == ""
    assert row_b["Availability"] == ""
    assert row_b["Room Name"] == ""
    assert row_b["Price"] == ""
    assert "N/A" not in row_b.values()
    assert "Unknown" not in row_b.values()


def test_flatten_rating_format_includes_label_and_count():
    rows = flatten_results(SAMPLE_RESULTS)
    row_a = next(r for r in rows if r["Hotel Name"].startswith("Hotel A"))
    assert row_a["Rating"] == "8.7 (Wonderful), 1234 reviews"
    assert row_a["Availability"] == "Only 1 left"
    assert row_a["Room Amenities"] == "TV; Air conditioning"
    assert row_a["Property Amenities"] == "Free Wifi; Pool"
    assert row_a["Image URL"] == "https://example.com/room.jpg"
    assert row_a["Room Size"] == "20 m²"


def test_psv_uses_pipe_delimiter_and_quotes_embedded_pipes():
    rows = flatten_results(SAMPLE_RESULTS)
    text = rows_to_psv(rows)
    lines = text.strip("\r\n").split("\r\n")
    header = lines[0].split("|")
    assert header == EXPORT_COLUMNS
    # "Hotel A | Test" contains a literal pipe -- must be quoted, not break columns
    assert '"Hotel A | Test"' in lines[1]
    # every data row must parse back to exactly len(EXPORT_COLUMNS) fields
    import csv
    reader = csv.reader(io.StringIO(text), delimiter="|")
    parsed_rows = list(reader)
    for r in parsed_rows:
        assert len(r) == len(EXPORT_COLUMNS)


def test_psv_is_utf8_safe_with_non_ascii():
    rows = flatten_results([
        {"hotel": {"hotel_name": "ملون فندق", "city": "Dubai"}, "units": []}
    ])
    text = rows_to_psv(rows)
    assert "ملون" in text


def test_xlsx_is_a_real_workbook_not_html():
    rows = flatten_results(SAMPLE_RESULTS)
    data = rows_to_xlsx_bytes(rows)
    assert data[:2] == b"PK"  # xlsx is a zip container, not HTML text
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == EXPORT_COLUMNS
    first_data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert "Hotel A" in first_data_row[0]


def test_xlsx_missing_fields_are_empty_cells():
    rows = flatten_results(SAMPLE_RESULTS)
    data = rows_to_xlsx_bytes(rows)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    rows_out = list(ws.iter_rows(min_row=2, values_only=True))
    row_b = next(r for r in rows_out if r[0] == "Hotel B")
    rating_idx = EXPORT_COLUMNS.index("Rating")
    assert row_b[rating_idx] is None
