"""Flatten scraped results into export rows, and render them as PSV or XLSX.

Takes the same JSON structure Storage writes (a list of
``{hotel, units, source_url, collected_at}`` dicts) and produces one row per
room/unit -- the natural grain for a spreadsheet. Only fields that actually
exist are included; nothing here invents a value for a missing field.
"""

from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Optional

EXPORT_COLUMNS = [
    "Hotel Name",
    "Property Type",
    "Rating",
    "Availability",
    "City",
    "Country",
    "Address",
    "Hotel URL",
    "Room Name",
    "Beds",
    "Room Size",
    "Plan",
    "Cancellation",
    "Price",
    "Currency",
    "Room Amenities",
    "Property Amenities",
    "Image URL",
]


def _blank(v: Any) -> Any:
    """Normalize missing values to '' for spreadsheet output (never a
    placeholder word like 'N/A' -- just an empty, honestly-blank cell)."""
    return "" if v is None else v


def _format_rating(hotel: Dict[str, Any]) -> str:
    score = hotel.get("rating_score")
    if score is None:
        return ""
    label = hotel.get("rating_label")
    count = hotel.get("review_count")
    text = str(score)
    if label:
        text += f" ({label})"
    if count is not None:
        text += f", {count} reviews"
    return text


def _format_size(size_m2: Optional[float]) -> str:
    if size_m2 is None:
        return ""
    return f"{size_m2:g} m²"


def _join_amenities(amenities: Optional[List[str]]) -> str:
    return "; ".join(a for a in (amenities or []) if a)


def flatten_results(results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """One row per room/unit; one blank-room row for a hotel with none."""
    rows: List[Dict[str, Any]] = []
    for item in results or []:
        hotel = item.get("hotel") or {}
        units = item.get("units") or []
        base = {
            "Hotel Name": _blank(hotel.get("hotel_name")),
            "Property Type": _blank(hotel.get("property_type")),
            "Rating": _format_rating(hotel),
            "Availability": _blank(hotel.get("availability")),
            "City": _blank(hotel.get("city")),
            "Country": _blank(hotel.get("country")),
            "Address": _blank(hotel.get("address")),
            "Hotel URL": _blank(hotel.get("booking_url")),
            "Property Amenities": _join_amenities(hotel.get("amenities")),
        }
        if not units:
            row = dict(base)
            for col in ("Room Name", "Beds", "Room Size", "Plan", "Cancellation",
                        "Price", "Currency", "Room Amenities", "Image URL"):
                row[col] = ""
            rows.append(row)
            continue
        for unit in units:
            images = unit.get("images") or []
            row = dict(base)
            row.update({
                "Room Name": _blank(unit.get("unit_name")),
                "Beds": _blank(unit.get("beds_text")),
                "Room Size": _format_size(unit.get("size_m2")),
                "Plan": _blank(unit.get("plan_name")),
                "Cancellation": _blank(unit.get("cancellation")),
                "Price": _blank(unit.get("price_amount")),
                "Currency": _blank(unit.get("price_currency")),
                "Room Amenities": _join_amenities(unit.get("amenities")),
                "Image URL": images[0] if images else "",
            })
            rows.append(row)
    return rows


def rows_to_psv(rows: List[Dict[str, Any]]) -> str:
    """Pipe-separated text, RFC-4180-style quoting via the stdlib csv module
    (delimiter='|') so a literal '|' or newline inside a cell (e.g. an
    amenity string) is properly escaped instead of corrupting columns."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter="|", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(EXPORT_COLUMNS)
    for row in rows:
        writer.writerow(["" if row.get(c) is None else row.get(c, "") for c in EXPORT_COLUMNS])
    return buf.getvalue()


def rows_to_xlsx_bytes(rows: List[Dict[str, Any]]) -> bytes:
    """A real .xlsx workbook (openpyxl), not an HTML table renamed."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Hotels"

    for col_idx, header in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col, "") or None)

    widths = {"Hotel Name": 30, "Address": 30, "Hotel URL": 40,
              "Room Amenities": 40, "Property Amenities": 40, "Image URL": 40}
    for col_idx, header in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 16)
    ws.freeze_panes = "A2"

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
